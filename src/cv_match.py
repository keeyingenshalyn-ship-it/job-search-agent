"""Score jobs against a CV.

Kept free of Streamlit (and pandas) so it can be unit tested on its own.

How the CV fit score (0-100) works — deliberately simple and explainable:
  * 70% skill coverage: MyCareersFuture tags every posting with a list of
    skills. We check how many of those skills your CV mentions (exact
    phrase, plus a few common aliases like "PowerBI" for "Power BI").
  * 30% text similarity: a TF-IDF cosine similarity between your CV and the
    job's title + skills + description, scaled so the best match in the run
    gets the full 30. This catches relevance the skill tags miss.
It's a ranking aid, not a verdict — read the matched/missing skills.
"""
from __future__ import annotations

import math
import re
from collections import Counter
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

SKILL_WEIGHT = 0.7
TEXT_WEIGHT = 0.3

# Different ways the same skill tends to be written. Keys are lowercase
# MyCareersFuture skill names; values are extra phrasings to look for in the CV.
SKILL_ALIASES: dict[str, list[str]] = {
    "power bi": ["powerbi", "power-bi"],
    "microsoft excel": ["excel"],
    "excel": ["microsoft excel"],
    "microsoft office": ["ms office", "excel", "powerpoint", "word"],
    "microsoft powerpoint": ["powerpoint"],
    "sql": ["mysql", "postgresql", "postgres", "t-sql", "sqlite", "bigquery"],
    "structured query language": ["sql"],
    "python": ["pandas", "numpy", "scikit-learn"],
    # Single-letter skills like "R" are only matched via these, never on
    # their own (a bare "r" shows up in too many unrelated places).
    "r": ["rstudio", "r studio", "r programming", "python/r", "r/python", "python r", "tidyverse", "ggplot2"],
    "tableau": ["tableau desktop", "tableau public"],
    "data visualization": ["data visualisation", "dashboard", "dashboards", "tableau", "power bi"],
    "data visualisation": ["data visualization", "dashboard", "dashboards", "tableau", "power bi"],
    "machine learning": ["predictive model", "predictive modelling", "predictive modeling", "regression", "classification"],
    "predictive modelling": ["predictive model", "predictive modeling", "machine learning"],
    "statistics": ["statistical", "regression"],
    "data analysis": ["data analytics", "analysed data", "analyzed data", "analysing data", "analyzing data"],
    "data analytics": ["data analysis", "analytics"],
    "business analytics": ["msc business analytics", "business analysis"],
    "stakeholder management": ["stakeholders", "stakeholder"],
    "mandarin": ["chinese", "mandarin chinese"],
    "chinese": ["mandarin"],
    "accounts receivable": ["ar ledger", "ar ledgers", "receivables"],
    "reconciliation": ["reconcile", "reconciled", "reconciliations"],
    "communication skills": ["communication"],
}

_STOPWORDS = set(
    """a an and are as at be by for from has have in is it its of on or that the to was
    were will with you your we our this these those i me my he she they them their
    not no but if so than then there here also into over under about after before
    such can may per via etc using use used work working job role team""".split()
)


# ---------------------------------------------------------------- reading CVs

def extract_text(path: str | Path) -> str:
    """Pull plain text out of a CV file (PDF, Word .docx, or plain text)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:  # pragma: no cover - depends on environment
            raise ImportError("Reading PDF CVs needs pypdf: pip install -r requirements.txt") from e
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError as e:  # pragma: no cover
            raise ImportError("Reading Word CVs needs python-docx: pip install -r requirements.txt") from e
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if ext in (".txt", ".md"):
        return path.read_text(errors="ignore")
    raise ValueError(f"Unsupported CV file type: {ext} (use PDF, .docx or .txt)")


def list_cvs(cv_dir: str | Path) -> list[Path]:
    cv_dir = Path(cv_dir)
    if not cv_dir.exists():
        return []
    return sorted(p for p in cv_dir.iterdir() if p.suffix.lower() in SUPPORTED_EXTENSIONS)


# ------------------------------------------------------------------- matching

def _normalize(text: str) -> str:
    # Pad with spaces so phrase checks can require word boundaries cheaply.
    text = text.lower().replace("’", "'")
    text = re.sub(r"[^a-z0-9+#.\-/' ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return f" {text} "


def _contains_phrase(norm_text: str, phrase: str) -> bool:
    # Normalize the phrase the same way as the CV ("Reporting & Analysis"
    # must match "reporting & analysis" after punctuation is stripped).
    phrase = _normalize(phrase).strip()
    if not phrase:
        return False
    # Word-boundary match so "r" doesn't match every word containing an r,
    # and "excel" doesn't match "excellent".
    pattern = r"(?<![a-z0-9])" + re.escape(phrase) + r"(?![a-z0-9])"
    return re.search(pattern, norm_text) is not None


def cv_has_skill(norm_cv: str, skill: str) -> bool:
    key = skill.lower().strip()
    if len(key) > 1 and _contains_phrase(norm_cv, key):
        return True
    return any(_contains_phrase(norm_cv, alias.strip()) for alias in SKILL_ALIASES.get(key, []))


def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"[a-z][a-z0-9+#]*", text.lower()) if t not in _STOPWORDS and len(t) > 1]


def _job_text(job: dict) -> str:
    title = job.get("title", "") or ""
    skills = " ".join(job.get("skills", []) or [])
    desc = job.get("description", "") or ""
    # Title repeated so it carries more weight than one line normally would.
    return f"{title} {title} {title} {skills} {desc}"


def _tfidf_vectors(docs: list[list[str]]) -> list[dict[str, float]]:
    n = len(docs)
    df = Counter()
    for d in docs:
        df.update(set(d))
    vectors = []
    for d in docs:
        tf = Counter(d)
        vec = {t: (1 + math.log(c)) * math.log((1 + n) / (1 + df[t]) + 1) for t, c in tf.items()}
        norm = math.sqrt(sum(v * v for v in vec.values())) or 1.0
        vectors.append({t: v / norm for t, v in vec.items()})
    return vectors


def _cosine(a: dict[str, float], b: dict[str, float]) -> float:
    if len(a) > len(b):
        a, b = b, a
    return sum(v * b.get(t, 0.0) for t, v in a.items())


def score_jobs(jobs: list[dict], cv_text: str) -> list[dict]:
    """Return one result per job (same order) with keys:
    cv_fit (0-100), matched_skills, missing_skills, skill_coverage (0-1)."""
    norm_cv = _normalize(cv_text)

    docs = [_tokens(cv_text)] + [_tokens(_job_text(j)) for j in jobs]
    vectors = _tfidf_vectors(docs)
    cv_vec, job_vecs = vectors[0], vectors[1:]
    sims = [_cosine(cv_vec, v) for v in job_vecs]
    max_sim = max(sims) if sims and max(sims) > 0 else 1.0

    results = []
    for job, sim in zip(jobs, sims):
        skills = [s for s in (job.get("skills") or []) if s]
        matched = [s for s in skills if cv_has_skill(norm_cv, s)]
        missing = [s for s in skills if s not in matched]
        coverage = len(matched) / len(skills) if skills else 0.0
        # A job with no skill tags is judged on text similarity alone.
        if skills:
            fit = 100 * (SKILL_WEIGHT * coverage + TEXT_WEIGHT * (sim / max_sim))
        else:
            fit = 100 * (sim / max_sim)
        results.append(
            {
                "cv_fit": round(fit, 1),
                "skill_coverage": coverage,
                "matched_skills": matched,
                "missing_skills": missing,
            }
        )
    return results


def cv_skill_summary(jobs: list[dict], cv_text: str, top: int = 15) -> tuple[list[str], list[str]]:
    """Across all jobs in a run: which commonly-requested skills the CV
    already shows, and which commonly-requested ones it doesn't mention.
    Each list is ordered by how many jobs ask for the skill."""
    norm_cv = _normalize(cv_text)
    demand = Counter()
    for j in jobs:
        demand.update({s for s in (j.get("skills") or []) if s})
    have, gaps = [], []
    for skill, _count in demand.most_common():
        (have if cv_has_skill(norm_cv, skill) else gaps).append(skill)
    return have[:top], gaps[:top]
